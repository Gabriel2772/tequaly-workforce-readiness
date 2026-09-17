import csv
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5_000


class ImportFileError(ValueError):
    pass


def parse_tabular_file(filename: str, content: bytes) -> tuple[list[str], list[dict[str, object]]]:
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError("file_too_large")
    suffix = Path(filename).suffix.casefold()
    if suffix == ".csv":
        return _parse_csv(content)
    if suffix == ".xlsx":
        return _parse_xlsx(content)
    raise ImportFileError("unsupported_file_format")


def _parse_csv(content: bytes) -> tuple[list[str], list[dict[str, object]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ImportFileError("invalid_utf8") from error
    reader = csv.DictReader(StringIO(text))
    headers = [header.strip() for header in (reader.fieldnames or []) if header is not None]
    if not headers:
        raise ImportFileError("missing_header")
    rows: list[dict[str, object]] = []
    for row in reader:
        if len(rows) >= MAX_ROWS:
            raise ImportFileError("row_limit_exceeded")
        rows.append({header: row.get(header, "") for header in headers})
    return headers, rows


def _parse_xlsx(content: bytes) -> tuple[list[str], list[dict[str, object]]]:
    try:
        with ZipFile(BytesIO(content)) as archive:
            if any(name.casefold().endswith("vbaproject.bin") for name in archive.namelist()):
                raise ImportFileError("macros_not_allowed")
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except (BadZipFile, KeyError, OSError, ValueError) as error:
        if isinstance(error, ImportFileError):
            raise
        raise ImportFileError("invalid_xlsx") from error
    worksheet = workbook.active
    if worksheet is None:
        workbook.close()
        raise ImportFileError("missing_worksheet")
    iterator = worksheet.iter_rows()
    header_cells = next(iterator, None)
    if header_cells is None:
        raise ImportFileError("missing_header")
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in header_cells]
    if not all(headers):
        raise ImportFileError("blank_header")
    rows: list[dict[str, object]] = []
    for cells in iterator:
        if len(rows) >= MAX_ROWS:
            raise ImportFileError("row_limit_exceeded")
        if any(cell.data_type == "f" for cell in cells):
            raise ImportFileError("formulas_not_allowed")
        values = [cell.value for cell in cells]
        if not any(value not in (None, "") for value in values):
            continue
        rows.append(
            {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
            }
        )
    workbook.close()
    return headers, rows
