import csv
import json
from datetime import UTC, datetime
from io import BytesIO, StringIO
from typing import cast
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.decisions.models import (
    DecisionOutcome,
    DecisionRun,
    DecisionSelection,
    DecisionTrainingAction,
)
from app.exports.types import ExportFormat, ExportType, StreamingExport
from app.operations.models import EligibilityResult, EligibilityRun, Operation, OperationRoleDemand
from app.workforce.models import Employee, Role, TrainingCatalog

MAX_EXPORT_ROWS = 10_000


def safe_spreadsheet_value(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


class ExportService:
    def __init__(self, session: Session, actor_id: str) -> None:
        self._session = session
        self._actor_id = actor_id

    def export(
        self,
        export_type: ExportType,
        export_format: ExportFormat,
        *,
        operation_id: UUID | None = None,
        active: bool | None = None,
    ) -> StreamingExport:
        headers, rows = self._dataset(export_type, operation_id=operation_id, active=active)
        if len(rows) > MAX_EXPORT_ROWS:
            raise ValueError("export_row_limit_exceeded")
        filters = {
            "operation_id": str(operation_id) if operation_id else "",
            "active": "" if active is None else str(active).lower(),
        }
        if export_format == "xlsx":
            content = self._xlsx(headers, rows, export_type, filters)
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            content = self._csv(headers, rows, export_type, filters)
            media_type = "text/csv; charset=utf-8"
        timestamp = datetime.now(UTC).strftime("%Y%m%d-%H%M%S")
        return StreamingExport(
            filename=f"twr-{export_type}-{timestamp}.{export_format}",
            media_type=media_type,
            content=content,
        )

    def _dataset(
        self,
        export_type: ExportType,
        *,
        operation_id: UUID | None,
        active: bool | None,
    ) -> tuple[list[str], list[list[object]]]:
        if export_type == "employees":
            return self._employees(active)
        if export_type in {"readiness", "gaps"}:
            return self._readiness(operation_id, gaps_only=export_type == "gaps")
        if export_type in {"scenarios", "decision_runs"}:
            return self._decisions(operation_id, include_audit=export_type == "decision_runs")
        if export_type == "training":
            return self._training(operation_id)
        return self._risk(operation_id)

    def _employees(self, active: bool | None) -> tuple[list[str], list[list[object]]]:
        query = select(Employee, Role).join(Role, Role.id == Employee.canonical_role_id)
        if active is not None:
            query = query.where(Employee.active.is_(active))
        records = self._session.execute(query.order_by(Employee.employee_number)).all()
        return [
            "matricula",
            "nome",
            "email",
            "cargo_codigo",
            "cargo",
            "base",
            "senioridade",
            "admissao",
            "ativo",
        ], [
            [
                employee.employee_number,
                employee.name,
                employee.email or "",
                role.code,
                role.name,
                employee.base_location,
                employee.seniority_level,
                employee.hired_on.isoformat() if employee.hired_on else "",
                employee.active,
            ]
            for employee, role in records
        ]

    def _readiness(
        self, operation_id: UUID | None, *, gaps_only: bool
    ) -> tuple[list[str], list[list[object]]]:
        query = (
            select(EligibilityResult, EligibilityRun, Employee, Operation, OperationRoleDemand)
            .join(EligibilityRun, EligibilityRun.id == EligibilityResult.eligibility_run_id)
            .join(Employee, Employee.id == EligibilityResult.employee_id)
            .join(Operation, Operation.id == EligibilityRun.operation_id)
            .join(OperationRoleDemand, OperationRoleDemand.id == EligibilityResult.role_demand_id)
        )
        if operation_id:
            query = query.where(Operation.id == operation_id)
        if gaps_only:
            query = query.where(EligibilityResult.classification != "ELIGIBLE")
        records = self._session.execute(
            query.order_by(Operation.code, Employee.employee_number)
        ).all()
        return [
            "operacao",
            "matricula",
            "nome",
            "demanda_id",
            "classificacao",
            "motivos",
            "gaps",
        ], [
            [
                operation.code,
                employee.employee_number,
                employee.name,
                str(demand.id),
                result.classification,
                json.dumps(result.reasons, ensure_ascii=False),
                json.dumps(result.gaps, ensure_ascii=False),
            ]
            for result, _run, employee, operation, demand in records
        ]

    def _decisions(
        self, operation_id: UUID | None, *, include_audit: bool
    ) -> tuple[list[str], list[list[object]]]:
        query = select(DecisionRun, Operation).join(
            Operation, Operation.id == DecisionRun.operation_id
        )
        if operation_id:
            query = query.where(Operation.id == operation_id)
        records = self._session.execute(query.order_by(DecisionRun.created_at.desc())).all()
        selections = {
            selection.decision_run_id: selection
            for selection in self._session.scalars(select(DecisionSelection))
        }
        outcomes = {
            outcome.decision_run_id: outcome
            for outcome in self._session.scalars(select(DecisionOutcome))
        }
        headers = [
            "decision_run_id",
            "operacao",
            "objetivo",
            "status",
            "solver",
            "regras",
            "hash_entrada",
            "custo_previsto_centavos",
            "criado_por",
            "criado_em",
        ]
        if include_audit:
            headers.extend(["selecionado_por", "selecionado_em", "resultado_registrado"])
        rows: list[list[object]] = []
        for run, operation in records:
            row: list[object] = [
                str(run.id),
                operation.code,
                run.objective,
                run.status,
                run.solver_version,
                run.rules_version,
                run.input_snapshot_hash,
                run.metrics.get("total_incremental_cost_cents", 0),
                run.created_by,
                run.created_at.isoformat(),
            ]
            if include_audit:
                selection = selections.get(run.id)
                row.extend(
                    [
                        selection.actor_id if selection else "",
                        selection.selected_at.isoformat() if selection else "",
                        run.id in outcomes,
                    ]
                )
            rows.append(row)
        return headers, rows

    def _training(self, operation_id: UUID | None) -> tuple[list[str], list[list[object]]]:
        query = (
            select(DecisionTrainingAction, DecisionRun, Employee, TrainingCatalog, Operation)
            .join(DecisionRun, DecisionRun.id == DecisionTrainingAction.decision_run_id)
            .join(Employee, Employee.id == DecisionTrainingAction.employee_id)
            .join(TrainingCatalog, TrainingCatalog.id == DecisionTrainingAction.training_catalog_id)
            .join(Operation, Operation.id == DecisionRun.operation_id)
        )
        if operation_id:
            query = query.where(Operation.id == operation_id)
        records = self._session.execute(query.order_by(Operation.code, Employee.name)).all()
        return [
            "operacao",
            "decision_run_id",
            "matricula",
            "nome",
            "treinamento",
            "pronto_em",
            "duracao_minutos",
            "custo_centavos",
        ], [
            [
                operation.code,
                str(run.id),
                employee.employee_number,
                employee.name,
                catalog.name,
                action.ready_at.isoformat(),
                action.duration_minutes,
                action.cost_cents,
            ]
            for action, run, employee, catalog, operation in records
        ]

    def _risk(self, operation_id: UUID | None) -> tuple[list[str], list[list[object]]]:
        from app.fragility.service import FragilityService

        report = FragilityService(self._session).calculate(
            3650, [operation_id] if operation_id else None
        )
        return [
            "operacao",
            "cargo",
            "turno",
            "exigido",
            "elegiveis",
            "treinaveis",
            "redundancia",
            "risco",
            "codigos",
        ], [
            [
                cell.operation_name,
                cell.role_name,
                cell.shift_code,
                cell.metric.required_count,
                cell.metric.eligible_count,
                cell.metric.trainable_count,
                cell.metric.redundancy,
                cell.risk.severity,
                ",".join(cell.risk.reason_codes),
            ]
            for cell in report.cells
        ]

    def _metadata(self, export_type: ExportType, filters: dict[str, str]) -> list[list[str]]:
        return [
            ["tipo", export_type],
            ["exportado_em_utc", datetime.now(UTC).isoformat()],
            ["ator", self._actor_id],
            ["versao_contrato", "1.0"],
            ["fuso", "UTC"],
            ["filtros", json.dumps(filters, ensure_ascii=False, sort_keys=True)],
        ]

    def _csv(
        self,
        headers: list[str],
        rows: list[list[object]],
        export_type: ExportType,
        filters: dict[str, str],
    ) -> bytes:
        output = StringIO(newline="")
        writer = csv.writer(output)
        for key, value in self._metadata(export_type, filters):
            writer.writerow([f"#{key}", value])
        writer.writerow([])
        writer.writerow(headers)
        writer.writerows([[safe_spreadsheet_value(value) for value in row] for row in rows])
        return output.getvalue().encode("utf-8-sig")

    def _xlsx(
        self,
        headers: list[str],
        rows: list[list[object]],
        export_type: ExportType,
        filters: dict[str, str],
    ) -> bytes:
        workbook = Workbook()
        data_sheet = workbook.active
        if data_sheet is None:
            raise RuntimeError("workbook_without_active_sheet")
        data_sheet.title = "Dados"
        data_sheet.append(headers)
        for row in rows:
            data_sheet.append([safe_spreadsheet_value(value) for value in row])
        data_sheet.freeze_panes = "A2"
        data_sheet.auto_filter.ref = data_sheet.dimensions
        header_fill = PatternFill("solid", fgColor="6C0775")
        for cell in data_sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
        metadata_sheet = workbook.create_sheet("Metadados")
        metadata_sheet.append(["Campo", "Valor"])
        for metadata_row in self._metadata(export_type, filters):
            metadata_sheet.append(metadata_row)
        for cell in metadata_sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()


def validate_export_type(value: str) -> ExportType:
    allowed = {
        "employees",
        "readiness",
        "gaps",
        "scenarios",
        "training",
        "risk",
        "decision_runs",
    }
    if value not in allowed:
        raise ValueError("unsupported_export_type")
    return cast(ExportType, value)
