from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.db.base import Base
from app.demo.generator import generate_demo_dataset
from app.demo.seed import seed_demo
from app.main import create_app
from app.workforce.models import Employee


def _client() -> TestClient:
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(engine, expire_on_commit=False)
    with factory() as session:
        seed_demo(session, generate_demo_dataset(seed=42, employee_count=20))
        employee = session.scalar(select(Employee).limit(1))
        assert employee is not None
        employee.name = '=HYPERLINK("https://example.invalid")'
        session.commit()
    return TestClient(create_app(session_factory=factory))


def test_filtered_csv_export_has_metadata_and_no_secrets() -> None:
    response = _client().get(
        "/exports/employees?format=csv&active=true",
        headers={"X-TWR-Actor": "viewer@example.com", "X-TWR-Role": "viewer"},
    )

    assert response.status_code == 200
    text = response.content.decode("utf-8-sig")
    assert "#ator,viewer@example.com" in text
    assert "matricula,nome,email" in text
    assert "'=HYPERLINK" in text
    assert "api_key" not in text.casefold()
    assert "token" not in text.casefold()
    assert "password" not in text.casefold()


def test_xlsx_export_has_data_and_metadata_sheets() -> None:
    response = _client().get("/exports/employees?format=xlsx&active=true")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=False)
    assert workbook.sheetnames == ["Dados", "Metadados"]
    assert workbook["Dados"]["A1"].value == "matricula"
    assert workbook["Dados"].max_row == 21
    assert workbook["Metadados"]["A2"].value == "tipo"
    workbook.close()


def test_unknown_export_type_is_rejected() -> None:
    response = _client().get("/exports/secrets")

    assert response.status_code == 422
